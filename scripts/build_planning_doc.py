#!/usr/bin/env python3
"""Rebuild the Google Docs importable technical plan using bundled python-docx.
Run the documents skill title sanitizer and render QA after generating this file.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
ROOT = Path(__file__).resolve().parents[1]
PAGES = [
('Hospital middleware development plan', [
('p', 'This document explains the implemented Agnos assignment for technical review. The delivery is a Go and Gin API with PostgreSQL, an HTTP hospital information system adapter, Nginx and Docker Compose. Staff authentication determines the hospital used for every patient search.'),
('h', 'Scope and approach'),
('p', 'The selected Approach C calls the configured hospital system when a national ID or passport ID is supplied, validates and stores the returned patient, and applies the search filters to that patient. Searches without an identifier use locally stored records. A synthetic HTTP mock makes the complete flow reproducible without real hospital credentials.'),
('p', 'The implementation covers staff creation, staff login and patient search with eight optional filters. The planning package includes the project structure, API contract and editable draw.io entity relationship diagram. It does not include a frontend, background synchronization or public production deployment.'),
('h', 'Review sequence'),
('table', ['Step', 'Review action'], [['1', 'Run make setup and make up to start the local stack.'], ['2', 'Run make test for Go vet, race-enabled tests and PostgreSQL integration coverage.'], ['3', 'Run make smoke to exercise requests through Nginx.'], ['4', 'Review docs/openapi.yaml and the editable ERD alongside the implementation.']]),
('h', 'Delivery boundaries'),
('p', 'The mock demonstrates the assumed patient shape over real HTTP. It cannot establish compatibility with a real HIS. Hospital B reuses the same mock shape to test hospital isolation. Real endpoint contracts, credentials, error responses and TLS requirements must be confirmed before an operational integration.'),
('p', 'Verification results and external GitHub and Google Docs destinations are tracked in docs/SUBMISSION.md. This plan describes the implementation and test strategy; the checklist is the source for executed validation evidence.')]),
('Project structure and responsibilities', [
('p', 'The code separates HTTP validation, application rules, persistence and external transport. Shared domain types allow the service to be tested without a live database or hospital system, while integration tests exercise the actual PostgreSQL repository.'),
('table', ['Path', 'Responsibility'], [
['cmd/api', 'Load configuration, connect PostgreSQL and serve the Gin API.'],
['cmd/mockhis', 'Serve synthetic hospital A and B responses over HTTP.'],
['internal/domain', 'Patient, staff, hospital, filters, errors and interfaces.'],
['internal/httpapi', 'Routes, strict JSON validation and status mapping.'],
['internal/service', 'Bcrypt, JWT authentication and search orchestration.'],
['internal/postgres', 'Parameterized queries, scoped upserts and pagination.'],
['internal/his', 'Timeouts, response validation and identifier checks.'],
['internal/config', 'Validated runtime configuration.'],
['migrations and seed', 'Schema and synthetic demonstration records.'],
['nginx and compose.yaml', 'Reverse proxy and local service topology.'],
['scripts', 'Bootstrap, verification, smoke and document utilities.'],
['docs and output', 'OpenAPI, planning, ERD and importable document.']]),
('h', 'Design rationale'),
('p', 'One service layer owns the search decision, keeping transport rules separate from database details. Hospital routing is configured on the server. Clients cannot provide a hospital ID or upstream URL to patient search. A repository interface supports focused unit tests without introducing a general framework.'),
('p', 'PostgreSQL is both the staff identity store and local patient store. A unique hospital and HN pair supports atomic refresh of a fetched patient. There is no queue or cache invalidation subsystem: identifier lookups always contact the HIS, and non-identifier searches explicitly cover local data only.')]),
('Request flow and hospital isolation', [
('h', 'Staff enrollment and login'),
('p', 'An operator provisions staff through POST /staff/create using the deployment registration key. The service resolves the hospital code, hashes the password with bcrypt and inserts the staff record. A username is unique within its hospital. Login finds the staff within the supplied hospital and verifies the password before issuing a one-hour HS256 JWT.'),
('h', 'Authenticated search'),
('table', ['Stage', 'Behavior'], [
['Authenticate', 'Validate JWT signature, issuer, audience and expiry; re-read the staff row by subject ID.'],
['Derive hospital', 'Use the hospital ID from the stored staff record; no client-supplied search hospital is accepted.'],
['Validate filters', 'Require a JSON object, exact field names, valid strings and bounded pagination.'],
['Choose source', 'National ID or passport ID uses HIS; all other requests use local PostgreSQL.'],
['Fetch and verify', 'Call configured HIS GET /patient/search/{id}; verify returned identifier, HN, date and gender.'],
['Persist and search', 'Upsert under the staff hospital; apply all filters with that hospital predicate. For HIS, restrict to the fetched local row ID.'],
['Return', 'Return a paginated data array, total matching count and source label.']]),
('h', 'Identifier and failure semantics'),
('p', 'When both identifiers are supplied, national ID selects the upstream request and passport ID remains an additional AND filter. A HIS 404 returns an empty result even if an older local row exists. An invalid response or network failure returns 502, an unconfigured HIS returns 503, and a timeout returns 504. There is no fallback to stale local data after a failed identifier lookup.'),
('p', 'The adapter ignores upstream local IDs and derives hospital ownership from the authenticated context. Redirects are refused. A three-second timeout is configured in Compose and the response limit is 1 MiB. Unknown upstream fields are tolerated; nonblank HN and the queried identifier are required.')]),
('Relational model', [
('p', 'Each hospital has zero or more staff and patients. Each staff or patient row belongs to exactly one hospital through a non-null foreign key. The diagram is generated from the implemented migration and is editable in draw.io.'),
('image', 'docs/diagrams/hospital-erd.png'),
('p', 'Editable source: docs/diagrams/hospital-erd.drawio. Vector preview: docs/diagrams/hospital-erd.svg. Schema authority: migrations/001_init.sql.'),
('h', 'Integrity and storage rules'),
('p', 'Hospitals use a unique code. Staff use UNIQUE (hospital_id, username); patients use UNIQUE (hospital_id, patient_hn). National and passport identifiers are not globally unique, allowing the same person in multiple hospitals. HN is required and cannot be blank. All twelve demographic columns may be null. Gender, when present, is M or F.'),
('p', 'The patient upsert refreshes all demographic fields, including null values, on a hospital and HN conflict. Indexed hospital and ID columns support stable paging; hospital and national or passport ID indexes support exact lookup. Page rows and total use one repeatable-read transaction.')]),
('Staff API contract', [
('p', 'All POST requests require Content-Type application/json and a single JSON object of at most 64 KiB. Unknown fields, duplicate fields, incorrectly cased keys, null required values, wrong types and trailing JSON are rejected with 400. Error responses use the common envelope described on page 7.'),
('h', 'Shared credentials body'),
('table', ['Field', 'Validation'], [
['username', 'Required string; trimmed; 3–64 ASCII letters, digits, dots, underscores or dashes. Case-sensitive.'],
['password', 'Required string; 8–72 UTF-8 bytes; preserved without trimming.'],
['hospital', 'Required string; trimmed; 1–64 ASCII letters, digits or dashes; must resolve to a configured hospital. Demo codes are A and B.']]),
('code', '{"username":"reviewer","password":"DemoPass123!",\n "hospital":"A"}'),
('h', 'POST staff create'),
('p', 'Path: /staff/create. Header: X-Registration-Key. A correct deployment key authorizes provisioning for either configured hospital. This is independent of the staff bearer token.'),
('code', '201  {"id":1,"username":"reviewer"}'),
('p', 'Expected errors: 400 for invalid input or unknown hospital; 403 for missing or incorrect registration key; 409 for an existing username within the hospital; 500 for internal or storage failure. The response omits the password hash and hospital ID.'),
('h', 'POST staff login'),
('p', 'Path: /staff/login. No provisioning key or bearer token is required. The same credentials validation applies. Wrong credentials and an unknown hospital both return 401; malformed input returns 400; storage failure returns 500.'),
('code', '200  {"access_token":"<JWT>",\n      "token_type":"Bearer","expires_in":3600}'),
('p', 'JWT uses HS256, issuer agnos, audience agnos-api and a staff ID subject. Expiry is required. Search requires Authorization: Bearer <JWT>. Removing the staff record invalidates subsequent authenticated requests because membership is re-read.')]),
('Patient search API contract', [
('p', 'POST /patient/search requires a valid staff bearer token. Every filter is optional. An empty object lists stored patients within the staff hospital. Explicit null, empty or whitespace-only filters are invalid. Strings are trimmed and limited to 1–255 UTF-8 bytes; identifiers have the tighter pattern below.'),
('table', ['Field', 'Match rule'], [
['national_id', 'Exact; 1–64 ASCII letters, digits or dashes. Triggers HIS lookup.'],
['passport_id', 'Exact; same identifier pattern. Triggers HIS when national_id is absent.'],
['first_name', 'Case-insensitive substring of first_name_th OR first_name_en.'],
['middle_name', 'Case-insensitive substring of middle_name_th OR middle_name_en.'],
['last_name', 'Case-insensitive substring of last_name_th OR last_name_en.'],
['date_of_birth', 'Exact real YYYY-MM-DD calendar date.'],
['phone_number', 'Exact string after trimming; no numeric conversion.'],
['email', 'Exact case-sensitive string after trimming.'],
['page', 'Optional integer 1–1000000; default 1.'],
['page_size', 'Optional integer 1–100; default 20.']]),
('p', 'All supplied filters combine with AND. Name patterns treat percent, underscore and backslash as literal text. Identifiers preserve leading zeroes. Search has no hospital, hospital_id or patient_hn input field. A name or other demographic filter never starts a broad remote HIS search.'),
('code', '{"first_name":"Somchai","date_of_birth":"1990-01-15",\n "page":1,"page_size":20}'),
('h', 'Pagination and source'),
('p', 'A successful response contains data (an array, including []), total (matching count before pagination), page, page_size and source. The source is local or his. Rows are ordered by the internal patient ID. A HIS search considers only the fetched row after filtering, so its total is zero or one. A page beyond the matches returns an empty data array while retaining the matching total.')]),
('Response fields and error behavior', [
('h', 'Patient object'),
('p', 'Every patient response contains the fields below. Demographic keys remain present with JSON null when no value is stored. Internal hospital ID is never serialized.'),
('table', ['Field group', 'JSON type and meaning'], [
['id', 'Integer int64; internal local patient row ID.'],
['patient_hn', 'Nonblank string; patient number within a hospital.'],
['first_name_th, middle_name_th, last_name_th', 'Each is string or null; Thai names.'],
['first_name_en, middle_name_en, last_name_en', 'Each is string or null; English names.'],
['date_of_birth', 'YYYY-MM-DD string or null.'],
['national_id, passport_id', 'Each is string or null; preserves leading zeroes.'],
['phone_number, email', 'Each is string or null.'],
['gender', 'M, F or null.']]),
('h', 'Application errors'),
('code', '{"error":{"code":"validation_error",\n          "message":"Null fields are not allowed"}}'),
('table', ['Status', 'Code and condition'], [
['400', 'validation_error — request validation; unknown hospital on create.'],
['401', 'unauthorized — login failure or invalid staff authentication.'],
['403', 'forbidden — registration key failure.'],
['409', 'conflict — duplicate username within one hospital.'],
['500', 'internal_error — internal or persistence failure.'],
['502', 'upstream_error — unreachable, malformed or mismatched HIS.'],
['503', 'his_unavailable — hospital system not configured.'],
['504', 'upstream_timeout — hospital request exceeded timeout.']]),
('p', 'The error object requires code and message and permits optional field context, currently unset. Responses avoid database details and upstream payloads. Nginx returns safe JSON errors, including 413 for oversized requests and 429 for staff rate limits. GET /healthz returns {"status":"ok"} for process liveness. GET /readyz pings PostgreSQL with a one-second deadline: 200 {"status":"ready"} or 503 {"status":"unavailable"}.')]),
('Verification and operational handoff', [
('h', 'Test strategy'),
('table', ['Layer', 'Positive and negative coverage'], [
['HTTP API', 'Create, login and search success; invalid credentials; missing auth; malformed, null, duplicate or unknown JSON; invalid pagination.'],
['Service', 'Password verification, token expiry and claims, hospital membership, source choice, upstream not-found and storage failures.'],
['HIS adapter', 'Real HTTP mock, both hospitals, identifier mismatch, malformed/oversized payload, invalid date/gender, redirects, timeout and unavailable configuration.'],
['PostgreSQL', 'Scoped uniqueness, all filters, literal wildcard handling, null refresh, stable pagination and overlapping cross-hospital identifiers.'],
['End to end', 'Live Nginx entry point, staff registration and login, local and HIS searches, authentication failures and hospital isolation.']]),
('p', 'Run make test for the containerized checks and make smoke against the running stack. scripts/check.sh executes go vet ./... and go test -race -count=1 ./.... A PostgreSQL test requires TEST_DATABASE_URL; the Compose test profile supplies it. Record actual outcomes in docs/SUBMISSION.md after completion.'),
('h', 'Local operation'),
('p', 'make setup generates .env with random secrets and restrictive file permissions, preserving existing configuration. make up builds and starts the services. Only Nginx binds a host port, on 127.0.0.1. The database volume persists when make down stops the stack; initialization SQL runs only for a new volume. Local Compose database transport uses sslmode=disable on its private network.'),
('h', 'Remaining integration work'),
('p', 'Before using a real HIS, verify each hospital contract and supply the confirmed base URLs, authentication and TLS settings. Local search coverage depends on seeded and previously fetched rows. Nginx limits staff creation and login to 10 requests per second per client IP with burst 20. There is no background synchronization, refresh token or production deployment. The provisioning key is a deployment-wide credential rather than a full administrator role system.'),
('p', 'Submit the repository and import output/agnos-development-plan.docx into Google Docs after the destinations are confirmed. Keep .env, tokens and real patient data out of source control. The synthetic demo does not require real patient information.')])
]

def make_table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    first = 0.6 if headers[0] == 'Step' else (2.65 if headers[0] == 'Field group' else (0.65 if headers[0] == 'Status' else 1.75))
    widths=[first,6.8-first]
    for i, col in enumerate(t.columns): col.width=Inches(widths[i])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row):cells[i].text=txt
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            cp.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(10);run.bold=ri==0
    props=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def main():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(25)
    doc.styles['Heading 1'].font.size=Pt(19)
    doc.styles['Heading 2'].font.size=Pt(13)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
    doc.core_properties.title='Hospital middleware development plan';doc.core_properties.author='';doc.core_properties.subject='Agnos implementation and API specification'
    footer=sec.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    md=['# Hospital middleware development plan','']
    for index,(title,blocks) in enumerate(PAGES):
        if index:doc.add_page_break()
        doc.add_paragraph(title,'Title' if index==0 else 'Heading 1')
        if index:md.extend(['## '+title,''])
        for block in blocks:
            kind=block[0]
            if kind=='h':doc.add_paragraph(block[1],'Heading 2');md.extend(['### '+block[1],''])
            elif kind=='p':doc.add_paragraph(block[1]);md.extend([block[1],''])
            elif kind=='code':
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
                r=p.add_run(block[1]);r.font.name='Menlo';r.font.size=Pt(10)
                md.extend(['```',block[1],'```',''])
            elif kind=='table':
                make_table(doc,block[1],block[2]);md.extend(['| '+' | '.join(block[1])+' |','| '+' | '.join(['---']*len(block[1]))+' |']+['| '+' | '.join(row)+' |' for row in block[2]]+[''])
            elif kind=='image':
                image=ROOT/block[1]
                if not image.exists():raise SystemExit('ERD PNG required: '+str(image))
                doc.add_picture(str(image),width=Inches(6.8));md.extend(['![Hospital ERD](diagrams/hospital-erd.svg)',''])
    (ROOT/'output').mkdir(exist_ok=True)
    doc.save(ROOT/'output/agnos-development-plan.docx')
    (ROOT/'docs/DEVELOPMENT_PLAN.md').write_text('\n'.join(md))
if __name__=='__main__':main()
